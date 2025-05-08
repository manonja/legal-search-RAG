"""DOCX document loader.

This module provides functionality to extract text from DOCX documents.
"""

import docx  # python-docx package
from pathlib import Path
from typing import Dict, Any

from app.core.struct_logger import log
from app.services.document_processor.loaders.loader_base import DocumentLoader


class DOCXLoader(DocumentLoader):
    """Loader for DOCX documents using python-docx."""

    def load(self, file_path: Path) -> tuple[str, Dict[str, Any]]:
        """Extract text from a DOCX file.

        Args:
            file_path: Path to the DOCX file

        Returns:
            Tuple containing extracted text and document metadata
        """
        text = ""
        metadata = {"original_file_type": "docx"}

        try:
            doc = docx.Document(file_path)

            # Extract document metadata from core properties
            core_props = doc.core_properties
            if core_props:
                for prop_name in dir(core_props):
                    if not prop_name.startswith("_"):
                        value = getattr(core_props, prop_name)
                        if value is not None:
                            metadata[f"docx_{prop_name}"] = str(value)

            # Extract text from paragraphs
            text = "\n".join(para.text for para in doc.paragraphs)

            # Add paragraph count to metadata
            metadata["paragraph_count"] = len(doc.paragraphs)

        except Exception as e:
            log.error("Error extracting DOCX", file_path=str(file_path), error=str(e))
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}") from e

        return text, metadata
